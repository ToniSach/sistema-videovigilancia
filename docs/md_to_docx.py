"""
Conversor Markdown -> Word (.docx) a medida para el documento de migracion.
Soporta: encabezados (#..######), tablas GFM (| ... |), listas (- / *),
citas (>), bloques de codigo (```), reglas (---), y formato inline
(**negrita**, `codigo`, [texto](enlace) -> texto).

Uso:
    python docs/md_to_docx.py docs/MIGRACION_STREAMING_WEBRTC.md docs/MIGRACION_STREAMING_WEBRTC.docx
"""
import re
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Pt, RGBColor, Inches

MONO = "Consolas"
CODE_BG = "F2F2F2"
HEADER_BG = "1F3864"  # azul oscuro para cabecera de tabla


def shade_cell(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


# ---------- parsing inline ----------
INLINE_RE = re.compile(
    r"(\*\*.+?\*\*)"          # **bold**
    r"|(`[^`]+`)"             # `code`
    r"|(\[[^\]]+\]\([^)]+\))" # [text](link)
)


def add_inline(paragraph, text):
    """Anade 'text' a 'paragraph' interpretando negrita, codigo y enlaces (solo el texto)."""
    pos = 0
    for m in INLINE_RE.finditer(text):
        if m.start() > pos:
            paragraph.add_run(text[pos:m.start()])
        token = m.group(0)
        if token.startswith("**"):
            r = paragraph.add_run(token[2:-2])
            r.bold = True
        elif token.startswith("`"):
            r = paragraph.add_run(token[1:-1])
            r.font.name = MONO
            r.font.size = Pt(9)
            r.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
        elif token.startswith("["):
            lm = re.match(r"\[([^\]]+)\]\(([^)]+)\)", token)
            label = lm.group(1)
            r = paragraph.add_run(label)
            r.font.color.rgb = RGBColor(0x1F, 0x49, 0xA0)
            r.underline = True
        pos = m.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def is_table_row(line):
    s = line.strip()
    return s.startswith("|") and s.endswith("|")


def is_separator_row(line):
    s = line.strip().strip("|")
    cells = [c.strip() for c in s.split("|")]
    return all(re.fullmatch(r":?-{2,}:?", c) for c in cells) and len(cells) > 0


def split_row(line):
    s = line.strip()
    if s.startswith("|"):
        s = s[1:]
    if s.endswith("|"):
        s = s[:-1]
    return [c.strip() for c in s.split("|")]


def main(src, dst):
    with open(src, encoding="utf-8") as f:
        lines = f.read().split("\n")

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)

    i = 0
    n = len(lines)
    while i < n:
        line = lines[i]
        stripped = line.strip()

        # ---- bloque de codigo ----
        if stripped.startswith("```"):
            i += 1
            code_lines = []
            while i < n and not lines[i].strip().startswith("```"):
                code_lines.append(lines[i])
                i += 1
            i += 1  # cerrar ```
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.2)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run("\n".join(code_lines) if code_lines else " ")
            run.font.name = MONO
            run.font.size = Pt(8.5)
            # sombreado del parrafo
            pPr = p._p.get_or_add_pPr()
            shd = OxmlElement("w:shd")
            shd.set(qn("w:val"), "clear")
            shd.set(qn("w:fill"), CODE_BG)
            pPr.append(shd)
            continue

        # ---- tabla ----
        if is_table_row(line) and (i + 1 < n) and is_separator_row(lines[i + 1]):
            header = split_row(line)
            i += 2  # saltar cabecera + separador
            rows = []
            while i < n and is_table_row(lines[i]):
                rows.append(split_row(lines[i]))
                i += 1
            ncols = len(header)
            table = doc.add_table(rows=1, cols=ncols)
            table.style = "Table Grid"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            hdr = table.rows[0].cells
            for c, htext in enumerate(header):
                shade_cell(hdr[c], HEADER_BG)
                para = hdr[c].paragraphs[0]
                add_inline(para, htext)
                for r in para.runs:
                    r.bold = True
                    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                    r.font.size = Pt(9)
            for row in rows:
                cells = table.add_row().cells
                for c in range(ncols):
                    txt = row[c] if c < len(row) else ""
                    para = cells[c].paragraphs[0]
                    add_inline(para, txt)
                    for r in para.runs:
                        r.font.size = Pt(8.5)
            doc.add_paragraph()
            continue

        # ---- regla horizontal ----
        if stripped == "---":
            doc.add_paragraph().add_run("").add_break()
            i += 1
            continue

        # ---- encabezados ----
        m = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if m:
            level = len(m.group(1))
            text = re.sub(r"[#*`]", "", m.group(2)).strip()
            # quitar enlaces dejando el texto
            text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)
            doc.add_heading(text, level=min(level, 4))
            i += 1
            continue

        # ---- cita ----
        if stripped.startswith(">"):
            text = stripped.lstrip(">").strip()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            add_inline(p, text)
            for r in p.runs:
                r.italic = True
                r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
            i += 1
            continue

        # ---- lista ----
        lm = re.match(r"^(\s*)[-*]\s+(.*)$", line)
        if lm:
            indent = len(lm.group(1))
            text = lm.group(2)
            stylename = "List Bullet" if indent < 2 else "List Bullet 2"
            try:
                p = doc.add_paragraph(style=stylename)
            except KeyError:
                p = doc.add_paragraph(style="List Bullet")
            add_inline(p, text)
            i += 1
            continue

        # ---- linea en blanco ----
        if stripped == "":
            i += 1
            continue

        # ---- parrafo normal ----
        p = doc.add_paragraph()
        add_inline(p, stripped)
        i += 1

    doc.save(dst)
    print(f"OK -> {dst}")


if __name__ == "__main__":
    src = sys.argv[1] if len(sys.argv) > 1 else "docs/MIGRACION_STREAMING_WEBRTC.md"
    dst = sys.argv[2] if len(sys.argv) > 2 else "docs/MIGRACION_STREAMING_WEBRTC.docx"
    main(src, dst)
